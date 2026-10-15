"""Small web dashboard for running the TradeIT crawler."""

from __future__ import annotations

import json
import re
import threading
import time
from io import BytesIO
from pathlib import Path
from typing import Any

import requests
from flask import Flask, jsonify, redirect, render_template, request, send_file, url_for
from openpyxl import Workbook, load_workbook

import merge_data
import tradeit_crawler as crawler


app = Flask(__name__)
job_lock = threading.Lock()
LEGACY_DATA_DIR = Path("data")
LEGACY_STATE_FILE = Path("state.json")
DATASETS_ROOT = Path("datasets")
LEGACY_DATASET_NAME = "legacy"
ACTIVE_THEME_KEY = "tradeit-theme"
job: dict[str, Any] = {
    "status": "idle",
    "message": "Sẵn sàng",
    "page": 0,
    "records": 0,
    "total": None,
    "account_index": 0,
    "account_total": 0,
    "account_label": "",
    "dataset_name": LEGACY_DATASET_NAME,
    "started_at": None,
}
stop_event = threading.Event()
ALLOWED_ACCOUNT_EXTENSIONS = {".xlsx", ".xlsm", ".xltx", ".xltm"}


def _page_path(raw_page: str) -> Path:
    try:
        page = int(raw_page)
    except (TypeError, ValueError) as exc:
        raise ValueError("Số trang không hợp lệ") from exc
    if page < 1:
        raise ValueError("Số trang phải lớn hơn 0")
    return crawler.page_file_path(page)


def _read_page(page: int) -> list[Any]:
    path = _page_path(str(page))
    if not path.is_file():
        raise FileNotFoundError(f"Không tìm thấy file trang {page}")
    payload = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload, list):
        raise ValueError(f"File trang {page} không chứa một mảng dữ liệu")
    return payload


def _dataset_page_path(dataset: str, page: int) -> Path:
    if page < 1:
        raise ValueError("Số trang phải lớn hơn 0")
    return _dataset_data_dir(dataset) / f"{page}.json"


def _read_page_in_dataset(dataset: str, page: int) -> list[Any]:
    path = _dataset_page_path(dataset, page)
    if not path.is_file():
        raise FileNotFoundError(f"Không tìm thấy file trang {page}")
    payload = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload, list):
        raise ValueError(f"File trang {page} không chứa một mảng dữ liệu")
    return payload


def _clean_name(name: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9_-]+", "_", name.strip()).strip("_-")
    if not slug:
        raise ValueError("Tên bộ dữ liệu không hợp lệ")
    return slug.lower()


def _dataset_dir(name: str) -> Path:
    if name == LEGACY_DATASET_NAME:
        return Path(".")
    return DATASETS_ROOT / _clean_name(name)


def _dataset_data_dir(name: str) -> Path:
    if name == LEGACY_DATASET_NAME:
        return LEGACY_DATA_DIR
    return _dataset_dir(name) / "data"


def _dataset_state_file(name: str) -> Path:
    if name == LEGACY_DATASET_NAME:
        return LEGACY_STATE_FILE
    return _dataset_dir(name) / "state.json"


def _dataset_csv_file(name: str) -> Path:
    if name == LEGACY_DATASET_NAME:
        return Path("data.csv")
    return _dataset_dir(name) / "data.csv"


def _dataset_display_name(name: str) -> str:
    if name == LEGACY_DATASET_NAME:
        return "Mặc định (legacy)"
    return name


def _dataset_exists(name: str) -> bool:
    data_dir = _dataset_data_dir(name)
    return data_dir.is_dir() and any(data_dir.glob("*.json"))


def _list_datasets() -> list[dict[str, Any]]:
    names = [LEGACY_DATASET_NAME]
    if DATASETS_ROOT.is_dir():
        for path in DATASETS_ROOT.iterdir():
            if path.is_dir():
                names.append(path.name)
    unique = []
    seen = set()
    for name in names:
        if name in seen:
            continue
        seen.add(name)
        unique.append(
            {
                "name": name,
                "label": _dataset_display_name(name),
                "files": len(list(_dataset_data_dir(name).glob("*.json")))
                if _dataset_data_dir(name).is_dir()
                else len(list(LEGACY_DATA_DIR.glob("*.json")))
                if name == LEGACY_DATASET_NAME and LEGACY_DATA_DIR.is_dir()
                else 0,
            }
        )
    return unique


def _default_dataset_name(importer: str, hs_code: str) -> str:
    return _clean_name(f"{importer}_{hs_code}")


def _ensure_dataset(name: str) -> str:
    dataset = _clean_name(name) if name != LEGACY_DATASET_NAME else LEGACY_DATASET_NAME
    if dataset != LEGACY_DATASET_NAME:
        _dataset_data_dir(dataset).mkdir(parents=True, exist_ok=True)
    return dataset


def _set_crawler_paths(name: str) -> str:
    dataset = _ensure_dataset(name)
    crawler.DATA_DIR = _dataset_data_dir(dataset)
    crawler.STATE_FILE = _dataset_state_file(dataset)
    return dataset


def _resolve_dataset_from_request(form: Any) -> str:
    dataset_name = str(form.get("dataset_name", "")).strip()
    if dataset_name:
        return dataset_name
    importer = str(form.get("importer", "netherland")).strip() or "netherland"
    hs_code = str(form.get("hs_code", "090411")).strip() or "090411"
    return _default_dataset_name(importer, hs_code)


def _fingerprint(item: Any) -> str:
    if isinstance(item, dict):
        normalized = merge_data.flatten(item)
    else:
        normalized = {"value": item}
    return json.dumps(normalized, ensure_ascii=False, sort_keys=True)


def _load_seen_fingerprints() -> set[str]:
    seen: set[str] = set()
    if not crawler.DATA_DIR.is_dir():
        return seen
    for path in crawler.DATA_DIR.glob("*.json"):
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            continue
        if not isinstance(payload, list):
            continue
        for item in payload:
            seen.add(_fingerprint(item))
    return seen


def _dedupe_items(items: list[Any], seen: set[str]) -> tuple[list[Any], int]:
    unique: list[Any] = []
    skipped = 0
    for item in items:
        key = _fingerprint(item)
        if key in seen:
            skipped += 1
            continue
        seen.add(key)
        unique.append(item)
    return unique, skipped


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _parse_accounts_excel(upload: Any) -> list[dict[str, str]]:
    """Parse an uploaded Excel file into token pairs.

    Expects the first sheet and two columns:
    - column A: token
    - column B: refresh token
    Header row is optional.
    """
    if upload is None or not getattr(upload, "filename", ""):
        raise ValueError("Vui lòng upload file Excel danh sách tài khoản")

    ext = Path(upload.filename).suffix.lower()
    if ext not in ALLOWED_ACCOUNT_EXTENSIONS:
        raise ValueError("Chỉ hỗ trợ file Excel .xlsx/.xlsm/.xltx/.xltm")

    workbook = load_workbook(upload.stream, read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    rows = sheet.iter_rows(values_only=True)

    try:
        first_row = next(rows)
    except StopIteration:
        workbook.close()
        raise ValueError("File Excel rỗng")

    header = [_cell_text(value).lower() for value in first_row[:4]]
    has_header = any("token" in value for value in header) and any(
        "refresh" in value or "rf" in value for value in header
    )

    if has_header:
        token_index = 0
        refresh_index = 1
        for idx, value in enumerate(header):
            if "token" in value and "refresh" not in value and "rf" not in value:
                token_index = idx
            if "refresh" in value or "rf" in value:
                refresh_index = idx
        data_rows = rows
    else:
        token_index = 0
        refresh_index = 1
        data_rows = iter([first_row, *rows])

    accounts: list[dict[str, str]] = []
    for row in data_rows:
        token = _cell_text(row[token_index] if len(row) > token_index else "")
        refresh = _cell_text(row[refresh_index] if len(row) > refresh_index else "")
        if token and refresh:
            accounts.append({"token": token, "refresh_token": refresh})

    workbook.close()
    return accounts


def _update(**values: Any) -> None:
    with job_lock:
        job.update(values)


def _run(options: dict[str, Any]) -> None:
    page = options["start_page"]
    total: int | None = None
    seen: set[str] = options["seen"]
    accounts: list[dict[str, str]] = options["accounts"]
    dataset_name: str = options["dataset_name"]
    _set_crawler_paths(dataset_name)
    body = {"importer": options["importer"], "hsCode": options["hs_code"]}

    if not options["force"]:
        saved_state = crawler.load_state()
        if saved_state:
            total = saved_state.get("total")

    _update(
        status="running",
        message="Đang chuẩn bị crawl",
        page=page,
        records=len(seen),
        total=total,
        account_index=0,
        account_total=len(accounts),
        account_label="",
        dataset_name=dataset_name,
    )

    try:
        if not accounts:
            raise ValueError("Chưa có tài khoản hợp lệ")

        for account_index, account in enumerate(accounts, start=1):
            if stop_event.is_set():
                break

            cookie = (
                f"{crawler.CRM_TOKEN_NAME}={account['token']}; "
                f"{crawler.CRM_RF_TOKEN_NAME}={account['refresh_token']}"
            )
            crawler.log_cookie_used(cookie)

            _update(
                account_index=account_index,
                account_total=len(accounts),
                account_label=f"Acc {account_index}/{len(accounts)}",
                page=page,
                records=len(seen),
                total=total,
                message=f"Đang crawl {account_index}/{len(accounts)} từ trang {page}",
            )

            session = requests.Session()
            try:
                while not stop_event.is_set():
                    if options["end_page"] is not None and page > options["end_page"]:
                        break

                    if not options["force"] and crawler.page_file_path(page).exists():
                        _update(
                            page=page,
                            records=len(seen),
                            message=(
                                f"Bỏ qua trang {page} đã có dữ liệu; "
                                f"tiếp tục {account_index}/{len(accounts)}"
                            ),
                        )
                        page += 1
                        continue

                    payload = fetch_page(session, cookie, page, options["per_page"], body)
                    data = payload.get("data") or {}
                    items = data.get("items")
                    if not isinstance(items, list):
                        raise RuntimeError(f"API không trả về data.items ở trang {page}")

                    row_count = data.get("rowCount", len(items))
                    if isinstance(data.get("total"), int):
                        total = data["total"]

                    unique_items, _skipped = _dedupe_items(items, seen)
                    crawler.save_page(unique_items, page, options["force"])

                    _update(
                        page=page,
                        records=len(seen),
                        total=total,
                        message=(
                            f"Tài khoản {account_index}/{len(accounts)}: "
                            f"trang {page} đã lưu {len(unique_items)} bản ghi mới"
                        ),
                    )
                    crawler.save_state(page, total, len(seen))

                    if isinstance(row_count, int) and row_count < options["per_page"]:
                        page += 1
                        break
                    if total is not None and len(seen) >= total:
                        page += 1
                        break

                    page += 1
                    stop_event.wait(options["delay"])
            except (crawler.AuthError, crawler.CreditError) as exc:
                if account_index >= len(accounts):
                    raise
                _update(
                    message=(
                        f"Tài khoản {account_index}/{len(accounts)} dừng ở trang {page}: "
                        f"{exc}. Chuyển sang tài khoản kế tiếp."
                    ),
                    records=len(seen),
                    total=total,
                )
            finally:
                session.close()

        if stop_event.is_set():
            _update(status="stopped", message="Đã dừng theo yêu cầu", records=len(seen))
        else:
            _update(status="completed", message=f"Hoàn tất: {len(seen)} bản ghi", records=len(seen))
    except (crawler.AuthError, crawler.CreditError) as exc:
        _update(
            status="error",
            message=f"Không crawl được tài khoản cuối ở trang {page}: {exc}.",
            records=len(seen),
            total=total,
        )
    except (requests.RequestException, RuntimeError, OSError, ValueError) as exc:
        _update(status="error", message=str(exc))


@app.get("/")
def index():
    with job_lock:
        current = dict(job)
    return render_template("index.html", job=current)


@app.post("/api/jobs")
def start_job():
    with job_lock:
        if job["status"] == "running":
            return jsonify({"error": "Đang có một job chạy"}), 409

    accounts_file = request.files.get("accounts_file")
    try:
        accounts = _parse_accounts_excel(accounts_file)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    if not accounts:
        return jsonify({"error": "Không đọc được tài khoản nào từ file Excel"}), 400

    try:
        start_page = max(1, int(request.form.get("start_page", "1")))
        end_raw = request.form.get("end_page", "").strip()
        end_page = int(end_raw) if end_raw else None
        per_page = min(50, max(1, int(request.form.get("per_page", "50"))))
        delay = max(0.0, float(request.form.get("delay", "1.5")))
    except ValueError:
        return jsonify({"error": "Thông số trang, số bản ghi hoặc delay không hợp lệ"}), 400

    if end_page is not None and end_page < start_page:
        return jsonify({"error": "Trang kết thúc phải lớn hơn hoặc bằng trang bắt đầu"}), 400

    dataset_raw = _resolve_dataset_from_request(request.form)
    dataset_name = _set_crawler_paths(dataset_raw)
    seen = _load_seen_fingerprints()
    options = {
        "start_page": start_page,
        "end_page": end_page,
        "per_page": per_page,
        "delay": delay,
        "force": request.form.get("force") == "on",
        "importer": request.form.get("importer", "netherland").strip() or "netherland",
        "hs_code": request.form.get("hs_code", "090411").strip() or "090411",
        "accounts": accounts,
        "seen": seen,
        "dataset_name": dataset_name,
    }

    stop_event.clear()
    with job_lock:
        job.update(
            status="running",
            message="Đang khởi động",
            page=start_page,
            records=len(seen),
            total=None,
            account_index=0,
            account_total=len(accounts),
            account_label="",
            dataset_name=dataset_name,
            started_at=time.time(),
        )

    threading.Thread(target=_run, args=(options,), daemon=True).start()
    return redirect(url_for("index"))


@app.post("/api/jobs/stop")
def stop_job():
    stop_event.set()
    return redirect(url_for("index"))


@app.get("/api/status")
def status():
    with job_lock:
        return jsonify(dict(job))


@app.get("/api/datasets")
def datasets():
    with job_lock:
        current = job.get("dataset_name", LEGACY_DATASET_NAME)
    return jsonify({"current": current, "datasets": _list_datasets()})


@app.get("/api/files")
def list_files():
    dataset_name = request.args.get("dataset", "").strip()
    if not dataset_name:
        with job_lock:
            dataset_name = str(job.get("dataset_name", LEGACY_DATASET_NAME))
    dataset_name = _ensure_dataset(dataset_name)
    data_dir = _dataset_data_dir(dataset_name)
    files = []
    if data_dir.is_dir():
        for path in sorted(data_dir.glob("*.json"), key=merge_data.page_number):
            page = merge_data.page_number(path)
            if page >= 10**9:
                continue
            try:
                count = len(_read_page_in_dataset(dataset_name, page))
                size = path.stat().st_size
                updated_at = path.stat().st_mtime
            except (OSError, ValueError, json.JSONDecodeError):
                count, size, updated_at = 0, path.stat().st_size, None
            files.append(
                {
                    "page": page,
                    "name": path.name,
                    "records": count,
                    "size": size,
                    "updated_at": updated_at,
                }
            )
    return jsonify(files)


@app.get("/api/files/<dataset>/<page>")
def get_file(dataset: str, page: str):
    try:
        dataset = _ensure_dataset(dataset)
        page_number = int(page)
        return jsonify({"page": page_number, "records": _read_page_in_dataset(dataset, page_number)})
    except (ValueError, FileNotFoundError, OSError, json.JSONDecodeError) as exc:
        return jsonify({"error": str(exc)}), 400


@app.put("/api/files/<dataset>/<page>")
def update_file(dataset: str, page: str):
    try:
        dataset = _ensure_dataset(dataset)
        page_number = int(page)
        path = _dataset_page_path(dataset, page_number)
        payload = request.get_json(silent=True)
        records = payload.get("records") if isinstance(payload, dict) else None
        if not isinstance(records, list):
            raise ValueError("Nội dung file phải là JSON array")
        json.dumps(records, ensure_ascii=False)
        temporary = path.with_suffix(".json.tmp")
        temporary.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
        temporary.replace(path)
        return jsonify({"page": page_number, "records": len(records), "message": "Đã lưu file"})
    except (ValueError, OSError, TypeError) as exc:
        return jsonify({"error": str(exc)}), 400


@app.post("/api/export")
def export_csv():
    selected = request.form.getlist("files")
    dataset_name = request.form.get("dataset", "").strip()
    if not dataset_name:
        with job_lock:
            dataset_name = str(job.get("dataset_name", LEGACY_DATASET_NAME))
    dataset_name = _ensure_dataset(dataset_name)
    data_dir = _dataset_data_dir(dataset_name)
    if selected:
        try:
            pages = sorted({int(value) for value in selected})
            records = []
            for page in pages:
                records.extend((page, merge_data.flatten(item)) for item in _read_page_in_dataset(dataset_name, page))
        except (ValueError, FileNotFoundError, OSError, json.JSONDecodeError) as exc:
            return jsonify({"error": str(exc)}), 400
    else:
        records = merge_data.load_records(data_dir)

    output = _dataset_csv_file(dataset_name)
    if not records:
        return jsonify({"error": "Chưa có dữ liệu để xuất"}), 400
    merge_data.write_csv(records, output, dedupe=True, encoding="utf-8-sig")
    return send_file(output, as_attachment=True, download_name=f"{dataset_name}-buyers.csv")


@app.get("/health")
def health():
    return jsonify({"status": "ok"})


@app.get("/download-template")
def download_template():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "accounts"
    sheet.append(["token", "refresh_token"])
    sheet.append(["token_1", "refresh_token_1"])
    sheet.append(["token_2", "refresh_token_2"])

    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    buffer.seek(0)
    return send_file(
        buffer,
        as_attachment=True,
        download_name="tradeit-accounts-template.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=False)
